"""
RapidEngine Credit Appraisal Engine — FastAPI Backend
======================================================
Endpoints cover:
  • Session / cookie-based user identity
  • Case creation + persistence (in-memory; swap Redis/Postgres for prod)
  • File upload + text extraction  (PyPDF2, openpyxl fallback)
  • Keyword/regex-based field parsing
  • Parsed-field save / load
  • Research agent  (DuckDuckGo via duckduckgo_search)
  • FinBERT sentiment analysis  (transformers pipeline)
  • Ollama LLM chat proxy
  • OpenAI-compatible LLM chat proxy
  • CAM report generation  (LLM-assisted, JSON-structured)
  • Immutable audit trail

Run:
    uvicorn main:app --reload --port 8000
"""

from __future__ import annotations

import io
import json
import re
import uuid
import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

import httpx
import openai
from duckduckgo_search import DDGS
from fastapi import (
    Cookie, FastAPI, File, Form, HTTPException,
    Request, Response, UploadFile,
)
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from transformers import pipeline
import PyPDF2

# ── Logging ───────────────────────────────────────────────────────────────────
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("rapidengine")

# ── App ───────────────────────────────────────────────────────────────────────
app = FastAPI(title="RapidEngine Credit Appraisal API", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],          # lock down in prod
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ── In-memory storage ─────────────────────────────────────────────────────────
USERS:     dict[str, dict] = {}   # session_id → user payload
CASES:     dict[str, dict] = {}   # case_id    → case payload
DOCUMENTS: dict[str, dict] = {}   # doc_id     → document payload
AUDIT_LOG: list[dict]       = []   # append-only

UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)

# ── FinBERT (lazy) ────────────────────────────────────────────────────────────
_finbert: Any = None

def get_finbert():
    global _finbert
    if _finbert is None:
        log.info("Loading FinBERT model…")
        _finbert = pipeline("text-classification", model="ProsusAI/finbert", top_k=None)
        log.info("FinBERT loaded.")
    return _finbert

# ──────────────────────────────────────────────────────────────────────────────
# KEYWORD / REGEX FIELD EXTRACTION
# Each entry: field_key → [list of regex patterns]
# First captured group (group 1) or full match is used as the value.
# ──────────────────────────────────────────────────────────────────────────────
FIELD_PATTERNS: dict[str, list[str]] = {
    # — Identity ——————————————————————————————————————————
    "company_name":      [r"(?:Taxpayer|Company Name|Account Name)[:\s]+([A-Za-z0-9 &.,()'-]+)",
                          r"^([A-Z][A-Za-z ]+(?:Ltd|Pvt|LLP|Inc|Limited))"],
    "gstin":             [r"GSTIN[:\s]+([0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z])"],
    "pan":               [r"PAN[:\s]+([A-Z]{5}[0-9]{4}[A-Z])"],
    "cin":               [r"CIN[:\s]+([A-Z][0-9]{5}[A-Z]{2}[0-9]{4}[A-Z]{3}[0-9]{6})"],
    "period":            [r"[Pp]eriod[:\s]+([A-Za-z]+ \d{4}\s*[–\-]\s*[A-Za-z]+ \d{4})"],

    # — GST ———————————————————————————————————————————————
    "annual_turnover":   [r"(?:TOTAL ANNUAL TURNOVER|Revenue from Operations)[^₹\n]*₹?([\d,]+(?:\.\d+)?(?:\s*Cr)?)"],
    "itc_claimed":       [r"(?:ITC Claimed|ITC in 3B)[:\s₹]*([\d,]+(?:\.\d+)?)"],
    "itc_2a":            [r"(?:ITC as per 2A|2A[^:]*:)[:\s₹]*([\d,]+(?:\.\d+)?)"],
    "itc_gap_pct":       [r"(?:DISCREPANCY|gap)[^%\n]*?(\d+(?:\.\d+)?)\s*%"],
    "total_tax_paid":    [r"(?:TOTAL TAX PAID|Tax Paid)[:\s₹]*([\d,]+(?:\.\d+)?)"],
    "filing_status":     [r"(All \d+ returns? filed[^\n]*)"],

    # — Bank ——————————————————————————————————————————————
    "bank_name":         [r"(?:^|\n)([A-Z]+ BANK)"],
    "account_number":    [r"Account No[:\s]+([X0-9 ]+)"],
    "total_credits":     [r"Total Credits?[:\s₹]*([\d,]+(?:\.\d+)?)"],
    "total_debits":      [r"Total Debits?[:\s₹]*([\d,]+(?:\.\d+)?)"],
    "closing_balance":   [r"Closing Balance[:\s₹]*([\d,]+(?:\.\d+)?)"],
    "cash_deposit_pct":  [r"Cash Deposits?[^%\n]*?(\d+(?:\.\d+)?)\s*%"],
    "cc_utilization":    [r"CC\s*[Ll]imit\s*[Uu]til[^₹\n]*₹?([\d,]+(?:\.\d+)?)"],
    "emi_detected":      [r"(?:Term Loan|TL)[^\n]*EMI[:\s₹]*([\d,]+(?:\.\d+)?)"],

    # — Financials (Annual Report / ITR) ——————————————————
    "revenue":           [r"Revenue from Operations[:\s₹]*([\d,]+(?:\.\d+)?(?:\s*Cr)?)"],
    "ebitda":            [r"EBITDA[:\s₹]*([\d,]+(?:\.\d+)?(?:\s*Cr)?)"],
    "ebitda_margin":     [r"EBITDA\s*[Mm]argin[:\s]*([\d.]+)\s*%"],
    "pat":               [r"PAT[:\s₹]*([\d,]+(?:\.\d+)?(?:\s*Cr)?)"],
    "net_worth":         [r"Net Worth[:\s₹]*([\d,]+(?:\.\d+)?(?:\s*Cr)?)"],
    "total_debt":        [r"Total Debt[:\s₹]*([\d,]+(?:\.\d+)?(?:\s*Cr)?)"],
    "debt_equity_ratio": [r"Debt[- ]Equity Ratio[:\s]*([\d.]+)\s*x?"],
    "dscr":              [r"DSCR[:\s]*([\d.]+)\s*x?"],
    "interest_coverage": [r"Interest Coverage[:\s]*([\d.]+)\s*x?"],
    "roc_charge":        [r"Charge\s*ID[^\n₹]*(₹[\d,]+(?:\.\d+)?(?:\s*Cr)?)"],
    "capacity_util":     [r"(\d+)\s*%\s*capacity"],
}

def extract_fields(text: str) -> dict[str, str]:
    """Run all FIELD_PATTERNS against raw text; return best match per key."""
    results: dict[str, str] = {}
    for field, patterns in FIELD_PATTERNS.items():
        for pat in patterns:
            m = re.search(pat, text, re.IGNORECASE | re.MULTILINE)
            if m:
                val = m.group(1).strip() if m.lastindex else m.group(0).strip()
                results[field] = val
                break
    return results

# ──────────────────────────────────────────────────────────────────────────────
# Helpers
# ──────────────────────────────────────────────────────────────────────────────
def new_id(prefix: str = "") -> str:
    return prefix + uuid.uuid4().hex[:12].upper()

def audit(case_id: str, event: str, actor: str, detail: str, severity: str = "info"):
    AUDIT_LOG.append({
        "ts":       datetime.utcnow().isoformat(),
        "case_id":  case_id,
        "event":    event,
        "actor":    actor,
        "detail":   detail,
        "severity": severity,
    })

def resolve_session(session_id: Optional[str]) -> str:
    if session_id and session_id in USERS:
        return session_id
    sid = new_id("SID")
    USERS[sid] = {
        "name": "R. Kumar",
        "role": "Senior Credit Officer",
        "created_at": datetime.utcnow().isoformat(),
    }
    return sid

def session_cases(sid: str) -> list[dict]:
    return [c for c in CASES.values() if c.get("session_id") == sid]

def case_docs(case_id: str) -> list[dict]:
    return [d for d in DOCUMENTS.values() if d.get("case_id") == case_id]

# ──────────────────────────────────────────────────────────────────────────────
# Text extraction
# ──────────────────────────────────────────────────────────────────────────────
def extract_text_pdf(data: bytes) -> str:
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(data))
        return "\n\n".join(
            p.extract_text() or "" for p in reader.pages
        )
    except Exception as e:
        log.warning("PyPDF2 error: %s", e)
        return ""

def extract_text(data: bytes, filename: str) -> str:
    fn = (filename or "").lower()
    if fn.endswith(".pdf"):
        return extract_text_pdf(data)
    if fn.endswith((".txt", ".csv", ".tsv", ".xml")):
        return data.decode("utf-8", errors="replace")
    if fn.endswith((".xlsx", ".xls")):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            lines: list[str] = []
            for ws in wb.worksheets:
                lines.append(f"=== Sheet: {ws.title} ===")
                for row in ws.iter_rows(values_only=True):
                    row_str = [str(c) if c is not None else "" for c in row]
                    if any(v.strip() for v in row_str):
                        lines.append("\t".join(row_str))
            return "\n".join(lines)
        except Exception as e:
            log.warning("openpyxl error: %s", e)
            return ""
    if fn.endswith(".docx"):
        try:
            import docx
            doc = docx.Document(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception as e:
            log.warning("python-docx error: %s", e)
    try:
        return data.decode("utf-8", errors="replace")
    except Exception:
        return "[Binary — extraction not supported]"

# ──────────────────────────────────────────────────────────────────────────────
# ══ 1. SESSION ════════════════════════════════════════════════════════════════
# ──────────────────────────────────────────────────────────────────────────────

@app.get("/api/session", tags=["Session"])
async def get_session(response: Response, session_id: Optional[str] = Cookie(None)):
    """Return or create a session. Sets httpOnly cookie."""
    sid = resolve_session(session_id)
    response.set_cookie("session_id", sid, max_age=86400 * 30, httponly=True, samesite="lax")
    return {"session_id": sid, "user": USERS[sid]}

# ──────────────────────────────────────────────────────────────────────────────
# ══ 2. CASES ══════════════════════════════════════════════════════════════════
# ──────────────────────────────────────────────────────────────────────────────

class NewCaseRequest(BaseModel):
    company_name: str
    gstin:        str   = ""
    cin:          str   = ""
    pan:          str   = ""
    sector:       str   = ""
    loan_amount:  float = 0.0
    purpose:      str   = ""

@app.post("/api/cases", tags=["Cases"])
async def create_case(
    payload:    NewCaseRequest,
    response:   Response,
    session_id: Optional[str] = Cookie(None),
):
    sid = resolve_session(session_id)
    response.set_cookie("session_id", sid, max_age=86400 * 30, httponly=True, samesite="lax")

    case_id = new_id("CASE")
    seq     = len(session_cases(sid)) + 1
    cam_ref = f"CAM-2025-{seq:03d}"

    case = {
        "case_id":               case_id,
        "cam_ref":               cam_ref,
        "session_id":            sid,
        "company_name":          payload.company_name,
        "gstin":                 payload.gstin,
        "cin":                   payload.cin,
        "pan":                   payload.pan,
        "sector":                payload.sector,
        "loan_amount":           payload.loan_amount,
        "purpose":               payload.purpose,
        "status":                "data_ingest",
        "score":                 None,
        "verdict":               None,
        "five_cs":               {},
        "conditions_of_sanction":[],
        "cam_sections":          {},
        "research_findings":     [],
        "officer_notes":         "",
        "created_at":            datetime.utcnow().isoformat(),
        "updated_at":            datetime.utcnow().isoformat(),
    }
    CASES[case_id] = case
    audit(case_id, "Case Created", USERS[sid]["name"],
          f"{payload.company_name} · ₹{payload.loan_amount} Cr requested")
    return case

@app.get("/api/cases", tags=["Cases"])
async def list_cases(session_id: Optional[str] = Cookie(None)):
    return session_cases(resolve_session(session_id))

@app.get("/api/cases/{case_id}", tags=["Cases"])
async def get_case(case_id: str, session_id: Optional[str] = Cookie(None)):
    sid  = resolve_session(session_id)
    case = CASES.get(case_id)
    if not case or case["session_id"] != sid:
        raise HTTPException(404, "Case not found")
    return case

@app.patch("/api/cases/{case_id}", tags=["Cases"])
async def update_case(
    case_id:    str,
    payload:    dict,
    session_id: Optional[str] = Cookie(None),
):
    sid  = resolve_session(session_id)
    case = CASES.get(case_id)
    if not case or case["session_id"] != sid:
        raise HTTPException(404, "Case not found")
    allowed = {"status", "officer_notes", "score", "verdict",
               "five_cs", "conditions_of_sanction"}
    for k, v in payload.items():
        if k in allowed:
            case[k] = v
    case["updated_at"] = datetime.utcnow().isoformat()
    audit(case_id, "Case Updated", USERS[sid]["name"], str({k: v for k, v in payload.items() if k in allowed}))
    return case

# ──────────────────────────────────────────────────────────────────────────────
# ══ 3. DOCUMENTS ══════════════════════════════════════════════════════════════
# ──────────────────────────────────────────────────────────────────────────────

@app.post("/api/cases/{case_id}/documents", tags=["Documents"])
async def upload_document(
    case_id:    str,
    file:       UploadFile          = File(...),
    doc_type:   str                 = Form("generic"),
    session_id: Optional[str]       = Cookie(None),
):
    """
    Upload a document, extract raw text, run regex parsing.
    Accepted doc_type values: gst | bank | annual | itr | board | rating | generic
    """
    sid  = resolve_session(session_id)
    case = CASES.get(case_id)
    if not case or case["session_id"] != sid:
        raise HTTPException(404, "Case not found")

    data     = await file.read()
    raw_text = extract_text(data, file.filename or "upload")
    parsed   = extract_fields(raw_text)

    # Simple confidence heuristic
    hits       = sum(1 for v in parsed.values() if v)
    confidence = min(98, max(30, int(hits / max(len(FIELD_PATTERNS), 1) * 100 * 2.5)))

    doc_id = new_id("DOC")
    dest   = UPLOAD_DIR / f"{doc_id}_{file.filename}"
    dest.write_bytes(data)

    doc = {
        "doc_id":        doc_id,
        "case_id":       case_id,
        "filename":      file.filename,
        "doc_type":      doc_type,
        "size_bytes":    len(data),
        "raw_text":      raw_text,
        "parsed_fields": parsed,
        "confidence":    confidence,
        "status":        "complete",
        "filepath":      str(dest),
        "uploaded_at":   datetime.utcnow().isoformat(),
    }
    DOCUMENTS[doc_id] = doc
    case["updated_at"] = datetime.utcnow().isoformat()
    audit(case_id, "Document Uploaded", USERS[sid]["name"],
          f"{file.filename} · {doc_type} · {confidence}% confidence",
          severity="info")

    # Return lightweight version
    return {**doc, "raw_text": raw_text[:400] + ("…" if len(raw_text) > 400 else "")}

@app.get("/api/cases/{case_id}/documents", tags=["Documents"])
async def list_documents(case_id: str, session_id: Optional[str] = Cookie(None)):
    sid  = resolve_session(session_id)
    case = CASES.get(case_id)
    if not case or case["session_id"] != sid:
        raise HTTPException(404, "Case not found")
    # Strip raw_text for listing (potentially large)
    return [{k: v for k, v in d.items() if k != "raw_text"} for d in case_docs(case_id)]

@app.get("/api/documents/{doc_id}", tags=["Documents"])
async def get_document(doc_id: str, session_id: Optional[str] = Cookie(None)):
    """Full document detail including raw_text — used by extraction modal."""
    sid = resolve_session(session_id)
    doc = DOCUMENTS.get(doc_id)
    if not doc:
        raise HTTPException(404, "Document not found")
    case = CASES.get(doc["case_id"])
    if not case or case["session_id"] != sid:
        raise HTTPException(403, "Access denied")
    return doc

@app.patch("/api/documents/{doc_id}/fields", tags=["Documents"])
async def save_fields(
    doc_id:     str,
    payload:    dict,
    session_id: Optional[str] = Cookie(None),
):
    """Persist user-edited parsed fields from the extraction modal."""
    sid = resolve_session(session_id)
    doc = DOCUMENTS.get(doc_id)
    if not doc:
        raise HTTPException(404, "Document not found")
    case = CASES.get(doc["case_id"])
    if not case or case["session_id"] != sid:
        raise HTTPException(403, "Access denied")
    doc["parsed_fields"].update(payload)
    doc["last_edited"] = datetime.utcnow().isoformat()
    audit(doc["case_id"], "Fields Edited", USERS[sid]["name"],
          f"{doc['filename']} — {len(payload)} field(s) saved")
    return doc["parsed_fields"]

# ──────────────────────────────────────────────────────────────────────────────
# ══ 4. RESEARCH AGENT ══════════════════════════════════════════════════════════
# ──────────────────────────────────────────────────────────────────────────────

class ResearchRequest(BaseModel):
    queries:                list[str]
    max_results_per_query:  int = 5

@app.post("/api/cases/{case_id}/research", tags=["Research"])
async def run_research(
    case_id:    str,
    payload:    ResearchRequest,
    session_id: Optional[str] = Cookie(None),
):
    """
    DuckDuckGo search for each query string.
    FinBERT sentiment applied to each headline.
    Results stored on the case.
    """
    sid  = resolve_session(session_id)
    case = CASES.get(case_id)
    if not case or case["session_id"] != sid:
        raise HTTPException(404, "Case not found")

    finbert  = get_finbert()
    findings: list[dict] = []
    ddgs = DDGS()

    for query in payload.queries[:4]:
        try:
            results = list(ddgs.text(query, max_results=payload.max_results_per_query))
        except Exception as e:
            log.warning("DDG search failed '%s': %s", query, e)
            results = []

        for r in results:
            title  = r.get("title", "")
            body   = r.get("body", "")
            url    = r.get("href", "")
            date   = r.get("published", "")
            snippet = (title + ". " + body[:200]).strip()

            sentiment_score = 0.0
            sentiment_label = "neutral"
            try:
                if snippet:
                    out    = finbert(snippet[:512])
                    scores = out[0] if isinstance(out[0], list) else out
                    best   = max(scores, key=lambda x: x["score"])
                    lbl    = best["label"].lower()
                    s      = best["score"]
                    sentiment_score = round(s if lbl == "positive" else (-s if lbl == "negative" else 0.0), 3)
                    sentiment_label = lbl
            except Exception as e:
                log.warning("FinBERT error: %s", e)

            findings.append({
                "query":           query,
                "title":           title,
                "body":            body[:400],
                "url":             url,
                "date":            date,
                "sentiment_label": sentiment_label,
                "sentiment_score": sentiment_score,
                "weight": (
                    "Critical" if sentiment_score < -0.6 else
                    "High"     if sentiment_score < -0.4 else
                    "Positive" if sentiment_score > 0.4  else "Neutral"
                ),
            })

    case["research_findings"] = findings
    case["updated_at"]        = datetime.utcnow().isoformat()

    avg = round(sum(f["sentiment_score"] for f in findings) / len(findings), 3) if findings else 0.0
    audit(case_id, "Research Complete", USERS[sid]["name"],
          f"{len(findings)} articles · avg sentiment {avg}")

    return _research_summary(findings)

@app.get("/api/cases/{case_id}/research", tags=["Research"])
async def get_research(case_id: str, session_id: Optional[str] = Cookie(None)):
    sid  = resolve_session(session_id)
    case = CASES.get(case_id)
    if not case or case["session_id"] != sid:
        raise HTTPException(404, "Case not found")
    return _research_summary(case.get("research_findings", []))

def _research_summary(findings: list[dict]) -> dict:
    avg = round(sum(f["sentiment_score"] for f in findings) / len(findings), 3) if findings else 0.0
    return {
        "findings":             findings,
        "aggregate_sentiment":  avg,
        "negative_count":       sum(1 for f in findings if f["sentiment_label"] == "negative"),
        "positive_count":       sum(1 for f in findings if f["sentiment_label"] == "positive"),
        "neutral_count":        sum(1 for f in findings if f["sentiment_label"] == "neutral"),
    }

@app.post("/api/sentiment", tags=["Research"])
async def sentiment_only(payload: dict):
    """Quick sentiment check on a list of texts (no case needed)."""
    texts   = payload.get("texts", [])[:20]
    finbert = get_finbert()
    results = []
    for text in texts:
        try:
            out    = finbert(text[:512])
            scores = out[0] if isinstance(out[0], list) else out
            best   = max(scores, key=lambda x: x["score"])
            lbl    = best["label"].lower()
            score  = round(best["score"] if lbl == "positive" else (-best["score"] if lbl == "negative" else 0.0), 3)
            results.append({"text": text[:80], "label": lbl, "score": score})
        except Exception as e:
            results.append({"text": text[:80], "label": "neutral", "score": 0.0, "error": str(e)})
    return {"results": results}

# ──────────────────────────────────────────────────────────────────────────────
# ══ 5. LLM — OLLAMA ═══════════════════════════════════════════════════════════
# ──────────────────────────────────────────────────────────────────────────────

class OllamaRequest(BaseModel):
    case_id:         str
    prompt:          str
    model:           str = "qwen2.5:72b"
    ollama_base_url: str = "http://localhost:11434"
    system_prompt:   str = (
        "You are a senior Indian credit analyst at a commercial bank. "
        "Provide concise, factual analysis grounded in RBI norms. "
        "Flag risks clearly. Use Indian currency and banking terminology."
    )

@app.post("/api/llm/ollama", tags=["LLM"])
async def ollama_chat(payload: OllamaRequest, session_id: Optional[str] = Cookie(None)):
    sid  = resolve_session(session_id)
    case = CASES.get(payload.case_id)
    if not case or case["session_id"] != sid:
        raise HTTPException(404, "Case not found")

    url  = f"{payload.ollama_base_url.rstrip('/')}/api/chat"
    body = {
        "model":   payload.model,
        "stream":  False,
        "messages":[
            {"role": "system", "content": payload.system_prompt},
            {"role": "user",   "content": payload.prompt},
        ],
        "options": {"temperature": 0.1, "num_ctx": 8192},
    }
    try:
        async with httpx.AsyncClient(timeout=120.0) as client:
            r = await client.post(url, json=body)
            r.raise_for_status()
        content = r.json().get("message", {}).get("content", "")
        audit(payload.case_id, "LLM Inference (Ollama)", USERS[sid]["name"],
              f"{payload.model} · {len(content)} chars")
        return {"model": payload.model, "provider": "ollama", "content": content}
    except httpx.ConnectError:
        raise HTTPException(503, f"Cannot reach Ollama at {payload.ollama_base_url}")
    except Exception as e:
        raise HTTPException(500, str(e))

@app.get("/api/ollama/models", tags=["LLM"])
async def list_ollama_models(base_url: str = "http://localhost:11434"):
    try:
        async with httpx.AsyncClient(timeout=5.0) as client:
            r = await client.get(f"{base_url.rstrip('/')}/api/tags")
            r.raise_for_status()
            return r.json()
    except Exception as e:
        return {"error": str(e), "models": []}

# ──────────────────────────────────────────────────────────────────────────────
# ══ 6. LLM — OPENAI COMPATIBLE ════════════════════════════════════════════════
# ──────────────────────────────────────────────────────────────────────────────

class OpenAIRequest(BaseModel):
    case_id:       str
    prompt:        str
    api_key:       str
    model:         str = "gpt-4o"
    base_url:      str = "https://api.openai.com/v1"
    system_prompt: str = (
        "You are a senior Indian credit analyst at a commercial bank. "
        "Provide concise, factual analysis grounded in RBI norms. "
        "Flag risks clearly. Use Indian currency and banking terminology."
    )

@app.post("/api/llm/openai", tags=["LLM"])
async def openai_chat(payload: OpenAIRequest, session_id: Optional[str] = Cookie(None)):
    sid  = resolve_session(session_id)
    case = CASES.get(payload.case_id)
    if not case or case["session_id"] != sid:
        raise HTTPException(404, "Case not found")
    try:
        client = openai.AsyncOpenAI(api_key=payload.api_key, base_url=payload.base_url)
        resp   = await client.chat.completions.create(
            model=payload.model,
            messages=[
                {"role": "system", "content": payload.system_prompt},
                {"role": "user",   "content": payload.prompt},
            ],
            temperature=0.1,
            max_tokens=2048,
        )
        content = resp.choices[0].message.content or ""
        audit(payload.case_id, "LLM Inference (OpenAI)", USERS[sid]["name"],
              f"{payload.model} · {len(content)} chars")
        return {"model": payload.model, "provider": "openai_compatible", "content": content}
    except openai.AuthenticationError:
        raise HTTPException(401, "Invalid API key")
    except openai.APIConnectionError:
        raise HTTPException(503, f"Cannot connect to {payload.base_url}")
    except Exception as e:
        raise HTTPException(500, str(e))

# ──────────────────────────────────────────────────────────────────────────────
# ══ 7. CAM GENERATION ══════════════════════════════════════════════════════════
# ──────────────────────────────────────────────────────────────────────────────

class CAMRequest(BaseModel):
    provider:         str = "ollama"                      # "ollama" | "openai"
    api_key:          str = ""                             # required for openai
    model:            str = "qwen2.5:72b"
    ollama_base_url:  str = "http://localhost:11434"
    openai_base_url:  str = "https://api.openai.com/v1"

def _build_cam_prompt(case: dict) -> str:
    docs   = case_docs(case["case_id"])
    merged: dict[str, str] = {}
    for d in docs:
        merged.update(d.get("parsed_fields", {}))

    research = case.get("research_findings", [])
    avg_sent = round(sum(r["sentiment_score"] for r in research) / len(research), 3) if research else 0.0
    neg      = sum(1 for r in research if r["sentiment_label"] == "negative")
    pos      = sum(1 for r in research if r["sentiment_label"] == "positive")
    headlines = "\n".join(
        f"  [{r['sentiment_score']:+.2f}] {r['title']}" for r in research[:8]
    )

    return f"""You are generating a Credit Appraisal Memo (CAM) for an Indian commercial bank.

CASE:
  Company      : {case['company_name']}
  GSTIN        : {case['gstin']}
  CIN          : {case['cin']}
  Sector       : {case['sector']}
  Loan Requested: ₹{case['loan_amount']} Cr
  Purpose      : {case['purpose']}
  CAM Ref      : {case['cam_ref']}

EXTRACTED FINANCIAL DATA:
{json.dumps(merged, indent=2)}

NEWS & RESEARCH ({len(research)} articles · avg sentiment {avg_sent}):
  Negative: {neg}  Positive: {pos}
{headlines}

INSTRUCTIONS:
Generate a structured CAM. Scores are 0–100.
Respond ONLY with valid JSON (no markdown fences, no preamble).

{{
  "five_cs": {{
    "character":  {{"score": <int>, "summary": "<2-3 sentences>"}},
    "capacity":   {{"score": <int>, "summary": "<2-3 sentences>"}},
    "capital":    {{"score": <int>, "summary": "<2-3 sentences>"}},
    "collateral": {{"score": <int>, "summary": "<2-3 sentences>"}},
    "conditions": {{"score": <int>, "summary": "<2-3 sentences>"}}
  }},
  "verdict": "APPROVE" | "CONDITIONAL APPROVE" | "REJECT",
  "recommended_rate": "<string>",
  "recommended_tenure": "<string>",
  "conditions_of_sanction": ["<str>", ...],
  "key_risks": ["<str>", ...],
  "executive_summary": "<3-4 sentences for the credit committee>"
}}""".strip()

def _weighted_score(five_cs: dict) -> float:
    weights = {"character": 0.20, "capacity": 0.30, "capital": 0.20,
               "collateral": 0.15, "conditions": 0.15}
    return round(sum(five_cs.get(k, {}).get("score", 50) * w for k, w in weights.items()), 2)

DEFAULT_FIVE_CS = {
    k: {"score": 60, "summary": "Insufficient data for automated assessment."}
    for k in ("character", "capacity", "capital", "collateral", "conditions")
}

@app.post("/api/cases/{case_id}/generate-cam", tags=["CAM"])
async def generate_cam(
    case_id:    str,
    payload:    CAMRequest,
    session_id: Optional[str] = Cookie(None),
):
    sid  = resolve_session(session_id)
    case = CASES.get(case_id)
    if not case or case["session_id"] != sid:
        raise HTTPException(404, "Case not found")

    prompt = _build_cam_prompt(case)

    # ── Call selected provider ──────────────────────────────
    if payload.provider == "openai":
        try:
            client = openai.AsyncOpenAI(api_key=payload.api_key, base_url=payload.openai_base_url)
            resp   = await client.chat.completions.create(
                model=payload.model,
                messages=[{"role": "user", "content": prompt}],
                temperature=0.05,
                max_tokens=2048,
                response_format={"type": "json_object"},
            )
            raw = resp.choices[0].message.content or "{}"
        except openai.AuthenticationError:
            raise HTTPException(401, "Invalid OpenAI API key")
        except Exception as e:
            raise HTTPException(500, f"OpenAI error: {e}")
    else:
        url  = f"{payload.ollama_base_url.rstrip('/')}/api/chat"
        body = {
            "model":   payload.model,
            "stream":  False,
            "messages":[{"role": "user", "content": prompt}],
            "options": {"temperature": 0.05, "num_ctx": 8192},
            "format":  "json",
        }
        try:
            async with httpx.AsyncClient(timeout=180.0) as client:
                r = await client.post(url, json=body)
                r.raise_for_status()
            raw = r.json().get("message", {}).get("content", "{}")
        except httpx.ConnectError:
            raise HTTPException(503, "Cannot reach Ollama. Is it running?")
        except Exception as e:
            raise HTTPException(500, str(e))

    # ── Parse JSON from LLM ─────────────────────────────────
    cam_data: dict = {}
    try:
        cam_data = json.loads(raw)
    except json.JSONDecodeError:
        m = re.search(r"\{.*\}", raw, re.DOTALL)
        if m:
            try:
                cam_data = json.loads(m.group(0))
            except Exception:
                pass

    five_cs = cam_data.get("five_cs") or DEFAULT_FIVE_CS
    score   = _weighted_score(five_cs)

    # ── Persist ─────────────────────────────────────────────
    case["five_cs"]                = five_cs
    case["score"]                  = score
    case["verdict"]                = cam_data.get("verdict", "CONDITIONAL APPROVE")
    case["conditions_of_sanction"] = cam_data.get("conditions_of_sanction", [])
    case["cam_sections"]           = {
        "executive_summary":  cam_data.get("executive_summary", ""),
        "key_risks":          cam_data.get("key_risks", []),
        "recommended_rate":   cam_data.get("recommended_rate", ""),
        "recommended_tenure": cam_data.get("recommended_tenure", ""),
    }
    case["status"]     = "cam_generated"
    case["updated_at"] = datetime.utcnow().isoformat()

    audit(case_id, "CAM Generated", USERS[sid]["name"],
          f"Score: {score} · Verdict: {case['verdict']}")

    return {**cam_data, "weighted_score": score}

@app.get("/api/cases/{case_id}/cam", tags=["CAM"])
async def get_cam(case_id: str, session_id: Optional[str] = Cookie(None)):
    sid  = resolve_session(session_id)
    case = CASES.get(case_id)
    if not case or case["session_id"] != sid:
        raise HTTPException(404, "Case not found")
    return {
        "case_id":               case_id,
        "cam_ref":               case.get("cam_ref"),
        "company_name":          case.get("company_name"),
        "loan_amount":           case.get("loan_amount"),
        "score":                 case.get("score"),
        "verdict":               case.get("verdict"),
        "five_cs":               case.get("five_cs", {}),
        "conditions_of_sanction":case.get("conditions_of_sanction", []),
        "cam_sections":          case.get("cam_sections", {}),
        "officer_notes":         case.get("officer_notes", ""),
    }

# ──────────────────────────────────────────────────────────────────────────────
# ══ 8. AUDIT TRAIL ═════════════════════════════════════════════════════════════
# ──────────────────────────────────────────────────────────────────────────────

@app.get("/api/audit", tags=["Audit"])
async def get_audit(
    case_id:    Optional[str] = None,
    session_id: Optional[str] = Cookie(None),
):
    sid      = resolve_session(session_id)
    my_cases = {c["case_id"] for c in session_cases(sid)}
    entries  = [e for e in AUDIT_LOG if e["case_id"] in my_cases]
    if case_id:
        entries = [e for e in entries if e["case_id"] == case_id]
    return sorted(entries, key=lambda x: x["ts"], reverse=True)

# ──────────────────────────────────────────────────────────────────────────────
# ══ 9. HEALTH ══════════════════════════════════════════════════════════════════
# ──────────────────────────────────────────────────────────────────────────────

@app.get("/api/health", tags=["Health"])
async def health():
    return {
        "status":        "ok",
        "cases":         len(CASES),
        "documents":     len(DOCUMENTS),
        "audit_entries": len(AUDIT_LOG),
        "ts":            datetime.utcnow().isoformat(),
    }

# ──────────────────────────────────────────────────────────────────────────────
# Entry point
# ──────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)